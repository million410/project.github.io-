# Main Libraries(TOOLS)
#-1-openpyxl------ # Best for: .xlsx files (modern Excel)
#-2-pandas-------- # Best for: Data analysis + tables

# Let's-- openpyxl
# -------Module Vs Class--------
#-Module-is simply a file (ending in .py) that contains code.
#--------It’s like a toolbox where you keep related tools together so you can find them later

from openpyxl import Workbook

# 1️⃣ Create & Write Excel File

# wb = Workbook()          #--Create workbook
# ws = wb.active           #--Select sheet
# ws.title = 'Sheet1'
#
# ws['A1'] = 'Name'
# ws['B1'] = 'Age'
# ws['A2'] = 'Million'
# ws['B2'] = 18
#
# wb.save('student.xlsx')

from openpyxl import load_workbook

# 2️⃣ Read Excel File

wb = load_workbook('student.xlsx')
ws = wb.active

print(ws['B2'].value)

# 3️⃣ Append Data (Add New Row)

# ws.append(['Ephrem', 19])
# ws.append(['Biruk', 20])
#
# wb.save('student.xlsx')

# 4️⃣ Modify / Update Cells

ws['B2'] = 19    # Change age
wb.save('student.xlsx')

# 5️⃣ Delete

# ws.delete_rows(5)
# wb.save('student.xlsx')

# ws.delete_cols(1)
# wb.save('student.xlsx')

# 6️⃣ Insert (Add Empty Row/Column)

# ws.insert_rows(2)    # insert empty row at 2
# ws.insert_cols(1)    # insert empty column at A

#-Goal                           code
# add a header row              ws.insert_rows(1)
# space b/n data                ws.insert_ros(5, 10)--(Adds 10 rows at row 5)
# shift data to column          ws.insert_cols(1, 2)--(Adds 2 columns at the start)

# ws.delete_rows(2)      # To remove the empty row at index 2
# ws.delete_cols(1)      # To remove the empty column at index 1 (Column A)
#
# wb.save('student.xlsx')

# 7️⃣ Move / Copy Data-----
#--Move range
# ws.move_range('A1:B2', rows=2, cols=1)
# wb.save('student.xlsx')

# 8️⃣ Create New Sheet
# wb.create_sheet('Sheet2')
# wb.save('student.xlsx')

# 9️⃣ Remove Sheet
# wb.remove(wb['Sheet1'])
# wb.save('student.xlsx')

# 🔟 Loop Through Data
# Rows---------
# for row in ws.iter_rows(values_only=True):
#     print(row)
#
# cell_list = []
# for row in ws.iter_rows(min_row=1, max_row=5, min_col=1, max_col=2):
#     for cell in row:
#         cell_list.append(cell.value)
# print(cell_list)

# Col---------------
# for col in ws.iter_cols(values_only=True):
#     print(col)
# for col in ws.iter_cols():
#     for cell in col:
#         print(cell.value)

# ⭐ Other Useful Things
# Get Max Rows & Columns

print(ws.max_row)
print(ws.max_column)

# Loop from row 1 to the very last row
for i in range(1, ws.max_row + 1):
    print(ws.cell(row=i, column=1).value)