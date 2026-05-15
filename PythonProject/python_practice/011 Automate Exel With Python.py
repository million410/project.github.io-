# Ready to exel your python skills

#DOCUMENTATION:
#----------------------
# https://openpyxl.readthedocs.io/en/stable/
# https://ssopenpyxl.readthedocs.io/en/stable/

import openpyxl
from openpyxl import load_workbook
filename  = 'Example.xlsx'
wb        = load_workbook(filename, data_only=True)
# print(wb)

from openpyxl import Workbook

# Get worksheets from workbook---------
#-----------------------------------------------
ws = wb.active
# all_ws = wb.worksheets
# all_ws_names = wb.sheetnames
#
# print(ws)
# print(all_ws)
# print(all_ws_names)

# for ws in all_ws:
#     print(ws.title)

#--Create/Copy/Remove Worksheets
#----------------------------------------------------------------
# new_ws = wb.create_sheet('NewSheet4') # Create a new Worksheet

# new_ws = wb.copy_worksheet(ws) # Copy worksheet
# new_ws.title = 'Copy!'         # Change ws name

# ws_to_remove = wb['NewSheet4']
# wb.remove(ws_to_remove)        # Remove worksheet
#
# wb.save(filename)

# 🔎 Read Data

# print(f'Worksheet: {ws.title}')
# print('-' * 50)
#
# #-------------------------------
# cell_a1 = ws['A1']
# cell_a2 = ws.cell(row=2, column=1)
#
# print(cell_a1.value)
# print(cell_a2.value)

# Loop over data (Range)

# for row in range(10):
#     row_data = []
#     for col in range(10):
#       cell = ws.cell(row=row + 1, column=col + 1)
#       if cell.value:
#          row_data.append(cell.value)
#     print(row_data)

# for row in ws.iter_rows():
#     row_data = [r.value for r in row]
#     print(row_data)

# Write Data ------------------------

# ws['B2'] = 'New B2 Value!'
# ws['B3']  = 'New B3 Value!'
# ws.cell(row=4, column=2, value='New B4 Value!')
#
# wb.save('Example.xlsx')
#
# # Append Data Row---------
# ws.append(['3--1', 'Mil', '18', 'America', '88'])
#
# wb.save('Example.xlsx')

# Insert/Delete Rows or Columns

# ws.insert_rows(3)
# ws.insert_cols(3)
#
# ws.delete_rows(1)
# ws.delete_cols(1)

# Move Data------

ws.move_range('A1:B10', rows=5, cols=5)
# Styling-----------------------
#--------------------------
