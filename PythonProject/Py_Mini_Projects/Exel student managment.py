from openpyxl import Workbook, load_workbook
import os
FILE = "Esm.xlsx"

if not os.path.exists(FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Students'
    ws.append(["ID", "Name", "Age", "Grade", "Average"])
    wb.save(FILE)

#--------------Load workbook-----------------
wb = load_workbook(FILE)
ws = wb["Students"]

# 1 Add Student
# 2 View Students
# 3 Search Student
# 4 Update Student
# 5 Delete Student
# 6 Exit
# """

#--------Add student----------
def add_student():
    sid = input('Inter Id: ')
    name = input('Name: ')
    age = input('Age: ')
    grade = input('Grade: ')
    average = input('Average: ')

    ws.append([sid, name, age, grade, average])
    wb.save(FILE)
    print('Successfully added!')
#----------View student--------
def view_students():
    for row in ws.iter_rows(min_row=2 ,values_only=True):
        print(row)
#---------search student--------
def search_student():
    sid = input('Input ID to search: ')
    for row in ws.iter_rows(min_row=2):
        if row[0].value == sid:
            print([cell.value for cell in row])
            return
    print('Not found')
#---------Update Student-----------------
def update_student():
    sid = input('Input ID to update: ')
    for row in ws.iter_rows(min_row=2):
        if row[0].value == sid:
            row[1].value = input('New Name:')
            row[2].value = int(input('New Age:'))
            row[3].value = input('New Grade:')
            row[4].value = float(input('New Average: '))
            wb.save(FILE)
            print('Updated!')
            return
    print('Not found')
#---------Delete----------------
def delete_student():
    sid = input('Input ID to delete: ')
    for i, row in enumerate(ws.iter_rows(min_row=2), start =  2):
        if row[0].value == sid:
            ws.delete_rows(i)
            wb.save(FILE)
            print('Deleted!') # row
            return
    print('Not found')
#-----------Menu------------------
while True:
    print('''
    1 Add Student
    2 View Students
    3 Search Students
    4 Update Student
    5 Delete Student
    6 Exit''')
    choice = input('Choose: ')
    if choice == '1':
        add_student()
    elif choice == '2':
        view_students()
    elif choice == '3':
        search_student()
    elif choice == '4':
        update_student()
    elif choice == '5':
        delete_student()
    elif choice == '6':
        break
    else:
        print('Invalid choice.')